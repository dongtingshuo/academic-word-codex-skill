#!/usr/bin/env python3
"""Build, analyze, and validate submission-ready academic DOCX files.

The script deliberately uses a conservative OOXML subset shared by Microsoft
Word, LibreOffice, and WPS. It never performs factual writing; the calling
agent supplies complete Markdown and structured citation data.
"""

from __future__ import annotations

import argparse
import copy
import json
import os
import posixpath
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

ALLOWED_FONTS = {"宋体", "黑体", "Times New Roman", "Arial", "Cambria Math"}
ALLOWED_FIELDS = {"TOC", "PAGE", "NUMPAGES", "SEQ", "REF", "PAGEREF", "STYLEREF", "HYPERLINK"}
PLACEHOLDER_RE = re.compile(
    r"(?:\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|\bTODO\b|待补(?:充|填写)?|目录将在打开文档时自动更新|X{3,}|x{5,})",
    re.IGNORECASE,
)
CITATION_RE = re.compile(r"\[@([A-Za-z0-9_.:-]+)\]")
CITATION_TOKEN_RE = re.compile(r"\[\[CITE:(\d+)\]\]")
VISIBLE_CITATION_RE = re.compile(r"\[(\d+)\]")
FIGURE_RE = re.compile(r"^\{\{figure:([A-Za-z0-9_.:-]+)\}\}$", re.IGNORECASE)
TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
SKILL_ROOT = Path(__file__).resolve().parents[1]
AUTHORITATIVE_DOTX = SKILL_ROOT / "assets" / "style-authority.dotx"


class BuildError(RuntimeError):
    pass


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as fh:
        value = json.load(fh)
    if not isinstance(value, dict):
        raise BuildError("Request must be a JSON object.")
    return value


def save_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(value, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def require_text(mapping: dict[str, Any], key: str, label: str | None = None) -> str:
    value = mapping.get(key)
    if not isinstance(value, str) or not value.strip():
        raise BuildError(f"Missing required value: {label or key}")
    return value.strip()


def validate_request(request: dict[str, Any]) -> None:
    mode = request.get("mode")
    doc_type = request.get("document_type")
    if mode not in {"draft", "format"}:
        raise BuildError("mode must be 'draft' or 'format'.")
    if doc_type not in {"course_report", "hust_thesis"}:
        raise BuildError("document_type must be 'course_report' or 'hust_thesis'.")

    metadata = request.get("metadata") or {}
    if not isinstance(metadata, dict):
        raise BuildError("metadata must be an object.")
    require_text(metadata, "title", "metadata.title")
    require_text(metadata, "author", "metadata.author")

    requirements = request.get("requirements") or {}
    if not isinstance(requirements, dict):
        raise BuildError("requirements must be an object.")
    for key in requirements.get("required_metadata", []):
        require_text(metadata, str(key), f"metadata.{key}")

    content = request.get("content") or {}
    if not isinstance(content, dict):
        raise BuildError("content must be an object.")
    if mode == "draft":
        markdown = content.get("markdown")
        markdown_path = content.get("markdown_path")
        if not (isinstance(markdown, str) and markdown.strip()) and not markdown_path:
            raise BuildError("draft mode requires content.markdown or content.markdown_path.")
    else:
        source = content.get("source_docx")
        if not source or not Path(source).is_file():
            raise BuildError("format mode requires a readable content.source_docx.")

    if doc_type == "hust_thesis":
        for key in (
            "student_id",
            "major",
            "advisor",
            "institution",
            "date",
            "title_en",
            "author_en",
            "major_en",
            "advisor_en",
            "degree_en",
            "date_en",
            "classification",
            "secrecy",
        ):
            require_text(metadata, key, f"metadata.{key}")
        require_text(content, "abstract_cn", "content.abstract_cn")
        require_text(content, "abstract_en", "content.abstract_en")
        if not content.get("keywords_cn") or not content.get("keywords_en"):
            raise BuildError("hust_thesis requires Chinese and English keyword lists.")

    external = requirements.get("external_template")
    if external and not Path(external).is_file():
        raise BuildError("requirements.external_template does not exist.")
    if external and not requirements.get("conflict_resolution"):
        raise BuildError("External template supplied but requirements.conflict_resolution is missing.")


def get_markdown(content: dict[str, Any]) -> str:
    markdown = content.get("markdown")
    if isinstance(markdown, str) and markdown.strip():
        return markdown.strip()
    path = Path(content["markdown_path"])
    if not path.is_file():
        raise BuildError(f"Markdown file does not exist: {path}")
    return path.read_text(encoding="utf-8").strip()


def set_east_asia_font(rpr: OxmlElement, name: str, latin: str) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), latin)


def style_font(style: Any, east_asia: str, latin: str, size_pt: float, bold: bool | None = None) -> None:
    style.font.name = latin
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    set_east_asia_font(rpr, east_asia, latin)


def run_font(run: Any, east_asia: str, latin: str, size_pt: float | None = None) -> None:
    run.font.name = latin
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    rpr = run._element.get_or_add_rPr()
    set_east_asia_font(rpr, east_asia, latin)


def get_or_add_style(doc: Document, name: str, style_type: WD_STYLE_TYPE, base: str | None = None) -> Any:
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, style_type)
    if base:
        try:
            style.base_style = doc.styles[base]
        except KeyError:
            pass
    return style


def remove_children(element: OxmlElement, tag: str) -> None:
    for child in list(element.findall(qn(tag))):
        element.remove(child)


def _style_index(styles: etree._Element) -> tuple[dict[str, str], dict[tuple[str, str], str]]:
    """Return styleId->name and (type, casefolded name)->styleId indexes."""
    by_id: dict[str, str] = {}
    by_name: dict[tuple[str, str], str] = {}
    for style in styles.xpath("//w:style", namespaces=NS):
        sid = style.get(qn("w:styleId"), "")
        stype = style.get(qn("w:type"), "")
        names = style.xpath("./w:name/@w:val", namespaces=NS)
        name = str(names[0]) if names else sid
        by_id[sid] = name
        by_name.setdefault((stype, name.casefold()), sid)
    return by_id, by_name


def apply_authoritative_template_parts(path: Path) -> None:
    """Apply the retained DOTX styles and numbering verbatim.

    The original Chinese template uses numeric styleIds (10/2/3/4) and links
    heading numbering through style-level numId 23. Rebuilding an equivalent
    English-ID numbering tree is not sufficient in Microsoft Word, especially
    while the default stylesWithEffects part remains attached. This gate maps
    every used style to the DOTX identifiers, removes direct heading numPr,
    copies the authoritative parts byte-for-byte, and removes stylesWithEffects.
    """
    if not AUTHORITATIVE_DOTX.is_file():
        raise BuildError(f"Authoritative DOTX is missing: {AUTHORITATIVE_DOTX}")

    with zipfile.ZipFile(path, "r") as generated, zipfile.ZipFile(AUTHORITATIVE_DOTX, "r") as template:
        infos = generated.infolist()
        package = {info.filename: generated.read(info.filename) for info in infos}
        generated_styles = etree.fromstring(package["word/styles.xml"])
        template_styles_bytes = template.read("word/styles.xml")
        template_styles = etree.fromstring(template_styles_bytes)
        story_names = [
            name for name in package
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            or name in {"word/footnotes.xml", "word/endnotes.xml"}
        ]
        stories = {name: etree.fromstring(package[name]) for name in story_names}
        document = stories["word/document.xml"]
        generated_by_id, _ = _style_index(generated_styles)
        _, template_by_name = _style_index(template_styles)

        aliases = {
            "题注": "caption",
        }
        heading_ids = {"heading 1": "10", "heading 2": "2", "heading 3": "3", "heading 4": "4"}

        for story_name, story in stories.items():
            for paragraph in story.xpath("//w:p", namespaces=NS):
                p_style_nodes = paragraph.xpath("./w:pPr/w:pStyle", namespaces=NS)
                if not p_style_nodes:
                    continue
                p_style = p_style_nodes[0]
                old_id = p_style.get(qn("w:val"), "")
                human_name = generated_by_id.get(old_id, old_id)
                folded = human_name.casefold()

                # The retained bibliography has no paragraph style: it uses a
                # 10 pt run size and a direct 360-DXA hanging indent.
                if story_name == "word/document.xml" and human_name == "参考文献":
                    ppr = paragraph.find("w:pPr", NS)
                    if ppr is not None:
                        ppr.remove(p_style)
                        remove_children(ppr, "w:ind")
                        ind = OxmlElement("w:ind")
                        ind.set(qn("w:left"), "360")
                        ind.set(qn("w:hanging"), "360")
                        ppr.append(ind)
                    for run in paragraph.xpath(".//w:r", namespaces=NS):
                        rpr = run.find("w:rPr", NS)
                        if rpr is None:
                            rpr = OxmlElement("w:rPr")
                            run.insert(0, rpr)
                        set_east_asia_font(rpr, "宋体", "Times New Roman")
                        remove_children(rpr, "w:sz")
                        remove_children(rpr, "w:szCs")
                        sz = OxmlElement("w:sz")
                        sz.set(qn("w:val"), "20")
                        sz_cs = OxmlElement("w:szCs")
                        sz_cs.set(qn("w:val"), "20")
                        rpr.extend([sz, sz_cs])
                    continue

                target_name = aliases.get(human_name, human_name)
                target_id = template_by_name.get(("paragraph", target_name.casefold()))
                if target_id:
                    p_style.set(qn("w:val"), target_id)

                # Original headings rely only on their numeric pStyle. Direct
                # numPr is deliberately absent in every template heading paragraph.
                if story_name == "word/document.xml" and folded in heading_ids:
                    p_style.set(qn("w:val"), heading_ids[folded])
                    ppr = paragraph.find("w:pPr", NS)
                    if ppr is not None:
                        remove_children(ppr, "w:numPr")

            # Map character and table styles in every story, including headers
            # and footers; otherwise Word sees dangling default-template IDs.
            for tag, stype in (("w:rStyle", "character"), ("w:tblStyle", "table")):
                for ref in story.xpath(f"//{tag}", namespaces=NS):
                    old_id = ref.get(qn("w:val"), "")
                    human_name = generated_by_id.get(old_id, old_id)
                    target_id = template_by_name.get((stype, aliases.get(human_name, human_name).casefold()))
                    if target_id:
                        ref.set(qn("w:val"), target_id)

        # The builder has only two non-heading numbering instances. Reuse the
        # DOTX's own bullet/decimal list definitions and its heading numId 23.
        for num_id in document.xpath("//w:pPr/w:numPr/w:numId", namespaces=NS):
            value = num_id.get(qn("w:val"), "")
            if value == "11":
                num_id.set(qn("w:val"), "33")
            elif value == "12":
                num_id.set(qn("w:val"), "32")

        for story_name, story in stories.items():
            package[story_name] = etree.tostring(
                story, xml_declaration=True, encoding="UTF-8", standalone=True
            )
        for part in (
            "word/styles.xml",
            "word/numbering.xml",
            "word/theme/theme1.xml",
            "word/fontTable.xml",
            "word/webSettings.xml",
        ):
            package[part] = template.read(part)

        # Preserve the template's punctuation compression, tab stop, language,
        # compatibility, math, and drawing defaults. The sole intentional
        # addition is automatic field refresh on open.
        settings = etree.fromstring(template.read("word/settings.xml"))
        for old in settings.xpath("./w:updateFields", namespaces=NS):
            settings.remove(old)
        update_fields = OxmlElement("w:updateFields")
        update_fields.set(qn("w:val"), "true")
        settings.append(update_fields)
        package["word/settings.xml"] = etree.tostring(
            settings, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        # settings.xml retains the template's separator references for IDs
        # -1 and 0. Microsoft Word requires the corresponding footnotes and
        # endnotes parts, relationships, and content-type overrides even when
        # the document contains no user notes. Omitting those companions makes
        # Word open the file with a "repaired Footnotes 1 / Endnotes 1" dialog.
        rels = etree.fromstring(package["word/_rels/document.xml.rels"])
        generated_rel_ids = {
            rel.get("Id", "") for rel in rels
        }
        numeric_rel_ids = [
            int(match.group(1))
            for rel_id in generated_rel_ids
            if (match := re.fullmatch(r"rId(\d+)", rel_id))
        ]
        next_rel_id = max(numeric_rel_ids, default=0) + 1
        template_rels = etree.fromstring(template.read("word/_rels/document.xml.rels"))
        content_types = etree.fromstring(package["[Content_Types].xml"])
        template_content_types = etree.fromstring(template.read("[Content_Types].xml"))
        for note_name, rel_suffix, content_type in (
            ("footnotes", "/footnotes", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"),
            ("endnotes", "/endnotes", "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"),
        ):
            part_name = f"word/{note_name}.xml"
            package[part_name] = template.read(part_name)
            if not any(str(rel.get("Type", "")).endswith(rel_suffix) for rel in rels):
                template_rel = next(
                    rel for rel in template_rels
                    if str(rel.get("Type", "")).endswith(rel_suffix)
                )
                new_rel = copy.deepcopy(template_rel)
                while f"rId{next_rel_id}" in generated_rel_ids:
                    next_rel_id += 1
                new_rel.set("Id", f"rId{next_rel_id}")
                generated_rel_ids.add(f"rId{next_rel_id}")
                next_rel_id += 1
                rels.append(new_rel)
            override_name = f"/{part_name}"
            if not any(node.get("PartName") == override_name for node in content_types):
                template_override = next(
                    (node for node in template_content_types if node.get("PartName") == override_name),
                    None,
                )
                if template_override is not None:
                    content_types.append(copy.deepcopy(template_override))
                else:
                    override = etree.Element("{http://schemas.openxmlformats.org/package/2006/content-types}Override")
                    override.set("PartName", override_name)
                    override.set("ContentType", content_type)
                    content_types.append(override)
        package["word/_rels/document.xml.rels"] = etree.tostring(
            rels, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        package["[Content_Types].xml"] = etree.tostring(
            content_types, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        # Word can prefer this parallel style part over styles.xml. The DOTX
        # does not contain it, so retaining python-docx's default copy causes
        # exactly the cross-application heading drift reported by the user.
        package.pop("word/stylesWithEffects.xml", None)
        rels = etree.fromstring(package["word/_rels/document.xml.rels"])
        for rel in list(rels):
            if str(rel.get("Type", "")).endswith("/stylesWithEffects"):
                rels.remove(rel)
        package["word/_rels/document.xml.rels"] = etree.tostring(
            rels, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        content_types = etree.fromstring(package["[Content_Types].xml"])
        for override in list(content_types):
            if override.get("PartName") == "/word/stylesWithEffects.xml":
                content_types.remove(override)
        package["[Content_Types].xml"] = etree.tostring(
            content_types, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    temp = path.with_name(f".{path.name}.template-parts.tmp")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as output_zip:
        written: set[str] = set()
        for info in infos:
            if info.filename not in package or info.filename in written:
                continue
            output_zip.writestr(info, package[info.filename])
            written.add(info.filename)
        for name, data in package.items():
            if name not in written:
                output_zip.writestr(name, data)
    temp.replace(path)


def append_spacing(ppr: OxmlElement, before: int, after: int, line: int | None = None) -> None:
    remove_children(ppr, "w:spacing")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    if line is not None:
        spacing.set(qn("w:line"), str(line))
        spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)


def configure_styles(doc: Document) -> dict[str, str]:
    normal = doc.styles["Normal"]
    style_font(normal, "宋体", "Times New Roman", 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    body = get_or_add_style(doc, "宋体小四", WD_STYLE_TYPE.PARAGRAPH, "Normal")
    style_font(body, "宋体", "Times New Roman", 12)
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.line_spacing = 1.25
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # These values are copied from the retained DOTX, including the different
    # Latin font on Heading 4 and the exact fixed line spacing at each level.
    heading_specs = [
        ("Heading 1", 16, "Arial", WD_ALIGN_PARAGRAPH.CENTER, 18, 18, 18, None, True),
        ("Heading 2", 14, "Arial", WD_ALIGN_PARAGRAPH.JUSTIFY, 10, 7, 18, 10.5, False),
        ("Heading 3", 12, "Arial", WD_ALIGN_PARAGRAPH.JUSTIFY, 7, 4, 16.1, None, False),
        ("Heading 4", 12, "黑体", WD_ALIGN_PARAGRAPH.JUSTIFY, 14, 14.5, 15, None, False),
    ]
    for level, (name, size, latin, align, before, after, line, right, page_break) in enumerate(heading_specs):
        style = doc.styles[name]
        style_font(style, "黑体", latin, size, False)
        style.font.italic = False
        pf = style.paragraph_format
        pf.alignment = align
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = Pt(line)
        pf.right_indent = Pt(right) if right is not None else None
        pf.keep_with_next = True
        pf.keep_together = True
        pf.page_break_before = page_break
        ppr = style.element.get_or_add_pPr()
        remove_children(ppr, "w:outlineLvl")
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level))
        ppr.append(outline)

    unnumbered = get_or_add_style(doc, "无编号标题", WD_STYLE_TYPE.PARAGRAPH, "Heading 1")
    style_font(unnumbered, "黑体", "Times New Roman", 16, True)
    unnumbered.font.italic = False
    unnumbered.paragraph_format.page_break_before = True
    unnumbered.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    unnumbered.paragraph_format.space_before = Pt(15.5)
    unnumbered.paragraph_format.space_after = Pt(14)
    ppr = unnumbered.element.get_or_add_pPr()
    remove_children(ppr, "w:numPr")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "0")
    num_pr.extend([ilvl, num_id])
    ppr.insert(0, num_pr)

    caption = get_or_add_style(doc, "题注", WD_STYLE_TYPE.PARAGRAPH, "Normal")
    style_font(caption, "黑体", "Arial", 10.5)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(7.6)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True

    figure = get_or_add_style(doc, "图片居中", WD_STYLE_TYPE.PARAGRAPH, "宋体小四")
    figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.first_line_indent = Pt(0)
    figure.paragraph_format.keep_with_next = True

    equation = get_or_add_style(doc, "公式", WD_STYLE_TYPE.PARAGRAPH, "Normal")
    style_font(equation, "宋体", "Cambria Math", 12)
    equation.paragraph_format.keep_together = True

    references = get_or_add_style(doc, "参考文献", WD_STYLE_TYPE.PARAGRAPH, "Normal")
    style_font(references, "宋体", "Times New Roman", 10.5)
    references.paragraph_format.left_indent = Pt(21)
    references.paragraph_format.first_line_indent = Pt(-21)
    references.paragraph_format.line_spacing = 1.2

    for toc_name, level in (("TOC 1", 1), ("TOC 2", 2), ("TOC 3", 3)):
        try:
            toc = doc.styles[toc_name]
        except KeyError:
            toc = doc.styles.add_style(toc_name, WD_STYLE_TYPE.PARAGRAPH)
        style_font(toc, "宋体" if level > 1 else "黑体", "Times New Roman", 12)

    # Materialize this built-in character style so page-number runs can be
    # remapped to the retained DOTX's `page number` style (styleId af2).
    get_or_add_style(doc, "Page Number", WD_STYLE_TYPE.CHARACTER, "Default Paragraph Font")

    return {
        "body": body.name,
        "caption": caption.name,
        "figure": figure.name,
        "equation": equation.name,
        "references": references.name,
        "unnumbered": unnumbered.name,
    }


def next_id(elements: Iterable[OxmlElement], attr: str) -> int:
    values = []
    for element in elements:
        value = element.get(qn(attr))
        if value is not None and str(value).isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def add_numbering(doc: Document) -> dict[str, int]:
    root = doc.part.numbering_part.element
    abs_id = next_id(root.findall(qn("w:abstractNum")), "w:abstractNumId")
    num_id = next_id(root.findall(qn("w:num")), "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)
    level_specs = (
        ("%1", "Arial", 32),
        ("%1.%2 ", "Arial", 28),
        ("%1.%2.%3 ", "Arial", 24),
        ("%1.%2.%3.%4  ", "黑体", 24),
    )
    for level, (text, latin, half_points) in enumerate(level_specs):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), doc.styles[f"Heading {level + 1}"].style_id)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        indent = 432 + level * 144
        ind.set(qn("w:left"), str(indent))
        ind.set(qn("w:hanging"), str(indent))
        ppr.append(ind)
        rpr = OxmlElement("w:rPr")
        set_east_asia_font(rpr, "黑体", latin)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), str(half_points))
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), str(half_points))
        rpr.extend([size, size_cs])
        lvl.extend([start, num_fmt, p_style, lvl_text, lvl_jc, ppr, rpr])
        abstract.append(lvl)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_ref = OxmlElement("w:abstractNumId")
    abs_ref.set(qn("w:val"), str(abs_id))
    num.append(abs_ref)
    root.append(num)
    heading_num_id = num_id

    def single_level(fmt: str, text: str) -> int:
        nonlocal abs_id, num_id
        abs_id += 1
        num_id += 1
        abstract_one = OxmlElement("w:abstractNum")
        abstract_one.set(qn("w:abstractNumId"), str(abs_id))
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        lvl.extend([start, num_fmt, lvl_text, ppr])
        abstract_one.append(lvl)
        root.append(abstract_one)
        num_one = OxmlElement("w:num")
        num_one.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num_one.append(ref)
        root.append(num_one)
        return num_id

    bullet_id = single_level("bullet", "•")
    decimal_id = single_level("decimal", "%1.")

    for level in range(4):
        style = doc.styles[f"Heading {level + 1}"]
        ppr = style.element.get_or_add_pPr()
        remove_children(ppr, "w:numPr")
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        num_ref = OxmlElement("w:numId")
        num_ref.set(qn("w:val"), str(heading_num_id))
        num_pr.extend([ilvl, num_ref])
        ppr.insert(0, num_pr)
    return {"headings": heading_num_id, "bullet": bullet_id, "decimal": decimal_id}


def set_paragraph_numbering(paragraph: Any, num_id: int, level: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    remove_children(ppr, "w:numPr")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    ppr.insert(0, num_pr)


def configure_page(section: Any, first_section: bool = False) -> None:
    """Apply the retained DOTX section geometry without unit-rounding drift."""
    section.page_width = Twips(11906)
    section.page_height = Twips(16838)
    section.top_margin = Twips(2552)
    section.right_margin = Twips(1588)
    section.bottom_margin = Twips(1588)
    section.left_margin = Twips(1588)
    section.header_distance = Twips(851)
    section.footer_distance = Twips(964)

    sect_pr = section._sectPr
    pg_sz = sect_pr.find(qn("w:pgSz"))
    if pg_sz is not None:
        pg_sz.set(qn("w:w"), "11906")
        pg_sz.set(qn("w:h"), "16838")
        if first_section:
            pg_sz.set(qn("w:code"), "9")
        else:
            pg_sz.attrib.pop(qn("w:code"), None)
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:space"), "425" if first_section else "720")
    grid = sect_pr.find(qn("w:docGrid"))
    if grid is None:
        grid = OxmlElement("w:docGrid")
        sect_pr.append(grid)
    grid.set(qn("w:type"), "linesAndChars")
    grid.set(qn("w:linePitch"), "312")


def set_page_numbering(section: Any, start: int = 1, fmt: str | None = "decimal") -> None:
    sect_pr = section._sectPr
    remove_children(sect_pr, "w:pgNumType")
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    if fmt:
        pg_num.set(qn("w:fmt"), fmt)
    sect_pr.append(pg_num)


def add_field(
    paragraph: Any,
    instruction: str,
    display: str = "",
    *,
    dirty: bool = False,
    locked: bool = False,
    hidden: bool = False,
) -> list[OxmlElement]:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    if locked:
        begin.set(qn("w:fldLock"), "true")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run._r.append(instr)

    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    result_run = paragraph.add_run(display)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    field_runs = [begin_run._r, instr_run._r, sep_run._r, result_run._r, end_run._r]
    if hidden:
        for field_run in field_runs:
            rpr = field_run.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                field_run.insert(0, rpr)
            rpr.append(OxmlElement("w:vanish"))
    return field_runs


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def clear_story(story: Any) -> None:
    element = story._element
    for child in list(element):
        element.remove(child)
    story.add_paragraph()


def set_paragraph_bottom_border(paragraph: Any, style: str = "single", size: int = 6) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        ppr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), style)
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)


def configure_header_footer(section: Any, header_text: str, page_number: bool = True) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_story(section.header)
    clear_story(section.footer)

    hp = section.header.paragraphs[0]
    hp.style = "Header"
    hp.add_run(header_text)

    fp = section.footer.paragraphs[0]
    fp.style = "Footer"
    if page_number:
        ppr = fp._p.get_or_add_pPr()
        frame = OxmlElement("w:framePr")
        frame.set(qn("w:wrap"), "around")
        frame.set(qn("w:vAnchor"), "text")
        frame.set(qn("w:hAnchor"), "margin")
        frame.set(qn("w:xAlign"), "center")
        frame.set(qn("w:y"), "1")
        ppr.append(frame)
        paragraph_rpr = OxmlElement("w:rPr")
        paragraph_rstyle = OxmlElement("w:rStyle")
        paragraph_rstyle.set(qn("w:val"), "PageNumber")
        paragraph_rpr.append(paragraph_rstyle)
        ppr.append(paragraph_rpr)
        add_field(fp, "PAGE", "1", dirty=True)
        for run in fp.runs:
            run.style = "Page Number"
        trailing = section.footer.add_paragraph(style="Footer")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")
        trailing._p.get_or_add_pPr().append(spacing)


def clear_header_footer(section: Any) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    clear_story(section.header)
    clear_story(section.footer)


def remove_table_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)


def add_centered_run(
    doc: Document,
    text: str,
    *,
    east_asia: str = "宋体",
    latin: str = "Times New Roman",
    size: float = 12,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
) -> Any:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    run_font(run, east_asia, latin, size)
    run.bold = bold
    return paragraph


def add_hust_english_title_page(doc: Document, metadata: dict[str, Any]) -> None:
    """Reproduce the retained template's second, English-language title page."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_centered_run(
        doc,
        "A Thesis Submitted in Partial Fulfillment of the Requirements",
        latin="Times New Roman",
        size=14,
        after=3,
    )
    add_centered_run(
        doc,
        f"for the Degree for the {metadata['degree_en']}",
        latin="Times New Roman",
        size=14,
    )
    add_centered_run(
        doc,
        str(metadata["title_en"]),
        latin="Times New Roman",
        size=18,
        bold=True,
        before=46,
        after=38,
    )
    add_centered_run(doc, f"Candidate: {metadata['author_en']}", latin="Times New Roman", size=14, after=8)
    add_centered_run(doc, f"Major: {metadata['major_en']}", latin="Times New Roman", size=14, after=8)
    add_centered_run(doc, f"Supervisor: {metadata['advisor_en']}", latin="Times New Roman", size=14)
    add_centered_run(
        doc,
        "Huazhong University of Science & Technology",
        latin="Times New Roman",
        size=14,
        before=50,
        after=3,
    )
    add_centered_run(doc, "Wuhan 430074, P.R.China", latin="Times New Roman", size=14, after=3)
    add_centered_run(doc, str(metadata["date_en"]), latin="Times New Roman", size=14)


def add_hust_declarations(doc: Document, metadata: dict[str, Any]) -> None:
    """Add the two declarations from the retained thesis template without placeholders."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_centered_run(doc, "独创性声明", east_asia="黑体", latin="Arial", size=16, bold=True, after=18)
    paragraph = doc.add_paragraph(style="宋体小四")
    paragraph.add_run(
        "本人声明所呈交的学位论文是我个人在导师指导下进行的研究工作及取得的研究成果。"
        "尽我所知，除文中已经标明引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写过的研究成果。"
        "对本文的研究做出贡献的个人和集体，均已在文中以明确方式标明。本人完全意识到本声明的法律结果由本人承担。"
    )
    signature = doc.add_paragraph(style="宋体小四")
    signature.paragraph_format.first_line_indent = None
    signature.add_run(f"学位论文作者签名：{metadata['author']}    日期：{metadata['date']}")

    add_centered_run(
        doc,
        "学位论文版权使用授权书",
        east_asia="黑体",
        latin="Arial",
        size=16,
        bold=True,
        before=28,
        after=18,
    )
    paragraph = doc.add_paragraph(style="宋体小四")
    paragraph.add_run(
        "本学位论文作者完全了解学校有关保留、使用学位论文的规定，即：学校有权保留并向国家有关部门或机构送交论文的复印件和电子版，"
        "允许论文被查阅和借阅。本人授权华中科技大学可以将本学位论文的全部或部分内容编入有关数据库进行检索，"
        "可以采用影印、缩印或扫描等复制手段保存和汇编本学位论文。"
    )
    secrecy = doc.add_paragraph(style="宋体小四")
    secrecy.paragraph_format.first_line_indent = None
    secrecy.add_run(f"本论文密级：{metadata['secrecy']}。")
    signatures = doc.add_paragraph(style="宋体小四")
    signatures.paragraph_format.first_line_indent = None
    signatures.add_run(
        f"学位论文作者签名：{metadata['author']}    指导教师签名：{metadata['advisor']}\n"
        f"日期：{metadata['date']}    日期：{metadata['date']}"
    )


def add_cover(doc: Document, metadata: dict[str, Any], doc_type: str) -> None:
    cover = doc.sections[0]
    configure_page(cover, first_section=True)
    # The retained template numbers the cover section internally with Roman
    # numerals but has no visible PAGE field there. The next section restarts
    # at I, so the cover remains visibly unnumbered.
    set_page_numbering(cover, 1, "upperRoman")
    clear_header_footer(cover)

    for _ in range(2):
        doc.add_paragraph()
    if doc_type == "hust_thesis":
        meta = doc.add_table(rows=2, cols=2)
        meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta.autofit = False
        remove_table_borders(meta)
        cover_meta = (
            (f"分类号  {metadata['classification']}", f"学号  {metadata['student_id']}"),
            ("学校代码  10487", f"密级  {metadata['secrecy']}"),
        )
        for row, values in zip(meta.rows, cover_meta):
            for cell, value in zip(row.cells, values):
                cell.width = Mm(75)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = paragraph.add_run(value)
                run_font(run, "宋体", "Times New Roman", 12)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(metadata.get("institution", "华中科技大学")))
        run_font(r, "华文中宋" if "华文中宋" in ALLOWED_FONTS else "黑体", "Arial", 24)
        r.bold = True
        p.space_after = Pt(12)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("硕士学位论文")
        run_font(r, "黑体", "Arial", 30)
        r.bold = True
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(metadata.get("institution", "课程报告")))
        run_font(r, "黑体", "Arial", 18)

    doc.add_paragraph()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(42)
    title = title_p.add_run(str(metadata["title"]))
    run_font(title, "黑体", "Arial", 22)
    title.bold = True

    fields: list[tuple[str, str]] = []
    if doc_type == "hust_thesis":
        fields = [
            ("学位申请人", str(metadata["author"])),
            ("学号", str(metadata["student_id"])),
            ("学科专业", str(metadata["major"])),
            ("指导教师", str(metadata["advisor"])),
            ("日期", str(metadata["date"])),
        ]
    else:
        candidates = [
            ("课程", metadata.get("course")),
            ("姓名", metadata.get("author")),
            ("学号", metadata.get("student_id")),
            ("教师", metadata.get("instructor")),
            ("日期", metadata.get("date")),
        ]
        fields = [(label, str(value)) for label, value in candidates if value]

    table = doc.add_table(rows=len(fields), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    for row, (label, value) in zip(table.rows, fields):
        row.cells[0].width = Mm(35)
        row.cells[1].width = Mm(75)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.DISTRIBUTE
        r0 = p0.add_run(label)
        run_font(r0, "黑体", "Arial", 12)
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(value)
        run_font(r1, "宋体", "Times New Roman", 12)

    if doc_type == "hust_thesis":
        add_hust_english_title_page(doc, metadata)
        add_hust_declarations(doc, metadata)


def markdown_outline(markdown: str) -> list[tuple[int, str, str]]:
    counters = [0, 0, 0, 0]
    entries: list[tuple[int, str, str]] = []
    for raw in markdown.splitlines():
        # The retained submission profile exposes only Heading 1 and Heading 2
        # in the table of contents.  Lower-level headings remain numbered in
        # the body but must never be materialized as cached TOC rows.
        match = re.match(r"^(#{1,2})\s+(.+?)\s*$", raw.strip())
        if not match:
            continue
        level = len(match.group(1))
        counters[level - 1] += 1
        for index in range(level, 4):
            counters[index] = 0
        number = ".".join(str(value) for value in counters[:level])
        entries.append((level, number, match.group(2).strip()))
    return entries


def format_cached_toc_paragraph(paragraph: Any, level: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    remove_children(ppr, "w:tabs")
    tabs = OxmlElement("w:tabs")
    if level > 1:
        left = OxmlElement("w:tab")
        left.set(qn("w:val"), "left")
        left.set(qn("w:pos"), str(840 * (level - 1)))
        tabs.append(left)
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:leader"), "middleDot")
    right.set(qn("w:pos"), "8720")
    tabs.append(right)
    ppr.append(tabs)
    for run in paragraph.runs:
        rpr = run._element.get_or_add_rPr()
        remove_children(rpr, "w:sz")
        remove_children(rpr, "w:szCs")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "21")
        no_proof = OxmlElement("w:noProof")
        # The retained template's cached TOC result fixes only East-Asian
        # text at 10.5 pt.  Latin/complex-script text (including 1.2) must
        # inherit 12 pt from the TOC/Normal style.  A direct szCs=22 made the
        # numeric prefix visibly too small.
        rpr.extend([no_proof, sz])


def add_toc(doc: Document, styles: dict[str, str], markdown: str) -> None:
    heading = doc.add_paragraph(style=styles["unnumbered"])
    heading.add_run("目录")
    entries = markdown_outline(markdown)
    if not entries:
        raise BuildError("TOC requested but Markdown contains no level 1-2 headings.")
    paragraphs = []
    for level, number, title in entries:
        paragraph = doc.add_paragraph(style=f"TOC {level}")
        prefix = "\t" if level > 1 else ""
        paragraph.add_run(f"{prefix}{number} {title}")
        format_cached_toc_paragraph(paragraph, level)
        paragraphs.append(paragraph)

    first = paragraphs[0]
    first_run = first.runs[0]._r
    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)
    instruction_run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-2" \\h \\z \\u '
    instruction_run.append(instruction)
    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    first._p.insert(first._p.index(first_run), begin_run)
    first._p.insert(first._p.index(first_run), instruction_run)
    first._p.insert(first._p.index(first_run), separate_run)
    end_run = paragraphs[-1].add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)


def cache_toc_pages(input_path: Path, page_map_path: Path, output_path: Path) -> None:
    page_map_raw = load_json(page_map_path)
    page_map = {str(key): str(value) for key, value in page_map_raw.items()}
    with zipfile.ZipFile(input_path, "r") as source:
        infos = source.infolist()
        package = {info.filename: source.read(info.filename) for info in infos}
    document = etree.fromstring(package["word/document.xml"])
    updated = 0
    for paragraph in document.xpath("//w:p[w:pPr/w:pStyle[@w:val='TOC1' or @w:val='TOC2']]", namespaces=NS):
        texts = [str(value) for value in paragraph.xpath(".//w:t/text()", namespaces=NS)]
        if not texts:
            continue
        visible = "".join(texts).strip()
        number_match = re.match(r"(\d+(?:\.\d+)*)\b", visible)
        number = number_match.group(1) if number_match else ""
        if number not in page_map:
            raise BuildError(f"TOC page map has no entry for heading {number!r}.")
        tab_run = OxmlElement("w:r")
        tab_run.append(OxmlElement("w:tab"))
        page_run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        no_proof = OxmlElement("w:noProof")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "21")
        rpr.extend([no_proof, sz])
        page_run.append(rpr)
        page_text = OxmlElement("w:t")
        page_text.text = page_map[number]
        page_run.append(page_text)
        end_runs = paragraph.xpath("./w:r[w:fldChar[@w:fldCharType='end']]", namespaces=NS)
        insert_at = paragraph.index(end_runs[0]) if end_runs else len(paragraph)
        paragraph.insert(insert_at, tab_run)
        paragraph.insert(insert_at + 1, page_run)
        updated += 1
    if updated != len(page_map):
        raise BuildError(f"TOC page map contains {len(page_map)} entries but updated {updated} cached entries.")
    package["word/document.xml"] = etree.tostring(
        document, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temp = output_path.with_name(f".{output_path.name}.toc-cache.tmp")
    with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as target:
        for info in infos:
            target.writestr(info, package[info.filename])
    temp.replace(output_path)


def add_keywords(doc: Document, label: str, keywords: Any, style_name: str) -> None:
    values = [str(item).strip() for item in keywords if str(item).strip()]
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.first_line_indent = Pt(0)
    lead = p.add_run(label)
    lead.bold = True
    p.add_run("；".join(values))


def set_cell_width(cell: Any, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table: Any, header_row: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for name, val, size in (
        ("top", "single", 12),
        ("bottom", "single", 12),
        ("left", "nil", 0),
        ("right", "nil", 0),
        ("insideV", "nil", 0),
        ("insideH", "nil", 0),
    ):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), val)
        if size:
            edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:color"), "000000")
        borders.append(edge)
    if header_row and table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        for cell in table.rows[0].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "000000")
            tc_borders.append(bottom)
            tc_pr.append(tc_borders)


def add_markdown_table(doc: Document, rows: list[list[str]], styles: dict[str, str]) -> None:
    if not rows or not rows[0]:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total_width = 8730
    base = total_width // cols
    widths = [base] * cols
    widths[-1] += total_width - sum(widths)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    set_table_borders(table)
    for ridx, row in enumerate(rows):
        for cidx in range(cols):
            cell = table.cell(ridx, cidx)
            set_cell_width(cell, widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[cidx] if cidx < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ridx == 0 or len(text) < 16 else WD_ALIGN_PARAGRAPH.LEFT
            add_text_with_citations(p, text, size_pt=10.5, bold=ridx == 0)


def add_seq_field(paragraph: Any, label: str, chapter: int, number: int) -> None:
    paragraph.add_run(f"{label} ")
    # This is the Word "Insert Caption / Include chapter number" structure:
    # label text + visible STYLEREF field + separator + visible SEQ field.
    # Lock only the cached STYLEREF result because LibreOffice incorrectly
    # expands it to the complete heading text; it remains a real Word field,
    # not typed numbering or a text box.
    add_field(paragraph, "STYLEREF 1 \\n", str(chapter), locked=True)
    paragraph.add_run("-")
    add_field(paragraph, f"SEQ {label} \\* ARABIC \\s 1", str(number))


def add_figure(doc: Document, asset: dict[str, Any], styles: dict[str, str], chapter: int, number: int) -> None:
    path = Path(str(asset.get("path", "")))
    if not path.is_file():
        raise BuildError(f"Figure asset does not exist: {path}")
    p = doc.add_paragraph(style=styles["figure"])
    run = p.add_run()
    width_mm = float(asset.get("width_mm", 140))
    run.add_picture(str(path), width=Mm(min(width_mm, 150)))
    drawing_props = p._p.xpath(".//*[local-name()='docPr']")
    if drawing_props:
        drawing_props[0].set("descr", str(asset.get("caption") or path.stem))
        drawing_props[0].set("title", str(asset.get("caption") or path.stem))
    caption = doc.add_paragraph(style=styles["caption"])
    add_seq_field(caption, "图", chapter, number)
    text = str(asset.get("caption", "")).strip()
    if text:
        caption.add_run(f" {text}")


def add_equation(doc: Document, equation_text: str, styles: dict[str, str], chapter: int, number: int) -> None:
    p = doc.add_paragraph(style=styles["equation"])
    ppr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    center = OxmlElement("w:tab")
    center.set(qn("w:val"), "center")
    center.set(qn("w:pos"), "4140")
    right = OxmlElement("w:tab")
    right.set(qn("w:val"), "right")
    right.set(qn("w:pos"), "8280")
    tabs.extend([center, right])
    ppr.append(tabs)
    p.add_run("\t")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    normalized_equation = equation_text.strip()
    for source, replacement in (
        ("\\\\times", "×"), ("\\times", "×"),
        ("\\\\cdot", "·"), ("\\cdot", "·"),
        ("\\\\leq", "≤"), ("\\leq", "≤"),
        ("\\\\geq", "≥"), ("\\geq", "≥"),
        ("\\\\neq", "≠"), ("\\neq", "≠"),
    ):
        normalized_equation = normalized_equation.replace(source, replacement)
    math_text.text = normalized_equation
    math_run.append(math_text)
    math.append(math_run)
    p._p.append(math)
    p.add_run("\t(")
    p.add_run(str(chapter))
    add_field(p, "STYLEREF 1 \\s", str(chapter), locked=True, hidden=True)
    p.add_run("-")
    add_field(p, "SEQ 公式 \\* ARABIC \\s 1", str(number))
    p.add_run(")")


def parse_table_row(line: str) -> list[str]:
    value = line.strip().strip("|")
    return [cell.strip() for cell in value.split("|")]


def citation_order(markdown: str, citations: list[dict[str, Any]]) -> tuple[str, list[dict[str, Any]]]:
    by_id: dict[str, dict[str, Any]] = {}
    for citation in citations:
        cid = str(citation.get("id", "")).strip()
        if not cid:
            raise BuildError("Every citation requires a non-empty id.")
        if cid in by_id:
            raise BuildError(f"Duplicate citation id: {cid}")
        by_id[cid] = citation
    order: list[str] = []
    for match in CITATION_RE.finditer(markdown):
        cid = match.group(1)
        if cid not in by_id:
            raise BuildError(f"Citation marker has no matching record: {cid}")
        if cid not in order:
            order.append(cid)
    number = {cid: idx + 1 for idx, cid in enumerate(order)}
    # Keep an internal token until paragraph runs are created so the visible
    # numeric citation can be emitted as a real superscript run.
    rendered = CITATION_RE.sub(lambda m: f"[[CITE:{number[m.group(1)]}]]", markdown)
    return rendered, [by_id[cid] for cid in order]


def add_text_with_citations(
    paragraph: Any,
    text: str,
    *,
    east_asia: str = "宋体",
    latin: str = "Times New Roman",
    size_pt: float | None = None,
    bold: bool = False,
) -> None:
    """Add text while rendering [[CITE:n]] as a superscript [n] run."""
    cursor = 0
    for match in CITATION_TOKEN_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            if size_pt is not None:
                run_font(run, east_asia, latin, size_pt)
            if bold:
                run.bold = True
        citation = paragraph.add_run(f"[{match.group(1)}]")
        rpr = citation._element.get_or_add_rPr()
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "080000")
        kern = OxmlElement("w:kern")
        kern.set(qn("w:val"), "0")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(round((size_pt or 12) * 2)))
        rpr.extend([color, kern, sz])
        citation.font.superscript = True
        if bold:
            citation.bold = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if size_pt is not None:
            run_font(run, east_asia, latin, size_pt)
        if bold:
            run.bold = True
    elif not text:
        paragraph.add_run("")


def format_citation(citation: dict[str, Any], number: int) -> str:
    if citation.get("rendered"):
        return f"[{number}] {str(citation['rendered']).strip()}"
    authors_raw = citation.get("authors") or []
    if isinstance(authors_raw, str):
        authors = authors_raw.strip()
    else:
        authors = ", ".join(str(item).strip() for item in authors_raw if str(item).strip())
    title = require_text(citation, "title", f"citation {citation.get('id')} title")
    year = str(citation.get("year", "")).strip()
    ctype = str(citation.get("type", "journal")).lower()
    prefix = f"[{number}] "
    author_part = f"{authors}. " if authors else ""
    if ctype == "journal":
        container = str(citation.get("container", "")).strip()
        if not container or not year:
            raise BuildError(f"Journal citation {citation.get('id')} requires container and year.")
        volume = str(citation.get("volume", "")).strip()
        issue = str(citation.get("issue", "")).strip()
        pages = str(citation.get("pages", "")).strip()
        vol_issue = volume + (f"({issue})" if issue else "")
        details = ", ".join(part for part in (year, vol_issue) if part)
        if pages:
            details += f": {pages}"
        text = f"{prefix}{author_part}{title}[J]. {container}, {details}."
    elif ctype == "book":
        publisher = str(citation.get("publisher", "")).strip()
        place = str(citation.get("place", "")).strip()
        if not publisher or not year:
            raise BuildError(f"Book citation {citation.get('id')} requires publisher and year.")
        pub = f"{place}: {publisher}" if place else publisher
        text = f"{prefix}{author_part}{title}[M]. {pub}, {year}."
    elif ctype in {"thesis", "dissertation"}:
        institution = str(citation.get("institution", "")).strip()
        place = str(citation.get("place", "")).strip()
        if not institution or not year:
            raise BuildError(f"Thesis citation {citation.get('id')} requires institution and year.")
        text = f"{prefix}{author_part}{title}[D]. {place + ': ' if place else ''}{institution}, {year}."
    elif ctype in {"web", "online"}:
        url = str(citation.get("url", "")).strip()
        accessed = str(citation.get("accessed", "")).strip()
        if not url or not accessed:
            raise BuildError(f"Online citation {citation.get('id')} requires url and accessed date.")
        date = f"({year})" if year else ""
        text = f"{prefix}{author_part}{title}[EB/OL]. {date}[{accessed}]. {url}."
    else:
        raise BuildError(f"Unsupported citation type for {citation.get('id')}: {ctype}")
    doi = str(citation.get("doi", "")).strip()
    if doi:
        text = text.rstrip(".") + f". DOI: {doi}."
    return text


def add_references(doc: Document, ordered: list[dict[str, Any]], styles: dict[str, str]) -> None:
    if not ordered:
        return
    heading = doc.add_paragraph("参考文献", style=styles["unnumbered"])
    heading.paragraph_format.page_break_before = True
    for idx, citation in enumerate(ordered, 1):
        doc.add_paragraph(format_citation(citation, idx), style=styles["references"])


def render_markdown(
    doc: Document,
    markdown: str,
    styles: dict[str, str],
    numbering: dict[str, int],
    assets: list[dict[str, Any]],
) -> None:
    asset_map = {str(asset.get("id")): asset for asset in assets}
    lines = markdown.splitlines()
    idx = 0
    chapter = 0
    figure_no = 0
    equation_no = 0
    while idx < len(lines):
        raw = lines[idx].rstrip()
        line = raw.strip()
        if not line:
            idx += 1
            continue
        if line == "---PAGEBREAK---":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            idx += 1
            continue

        heading_match = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            if level == 1:
                chapter += 1
                figure_no = 0
                equation_no = 0
            p = doc.add_paragraph(text, style=f"Heading {level}")
            set_paragraph_numbering(p, numbering["headings"], level - 1)
            idx += 1
            continue

        figure_match = FIGURE_RE.match(line)
        if figure_match:
            aid = figure_match.group(1)
            if aid not in asset_map:
                raise BuildError(f"Figure marker has no matching asset: {aid}")
            figure_no += 1
            add_figure(doc, asset_map[aid], styles, max(chapter, 1), figure_no)
            idx += 1
            continue

        if line.startswith("$$"):
            equation = line[2:]
            while not equation.rstrip().endswith("$$") and idx + 1 < len(lines):
                idx += 1
                equation += " " + lines[idx].strip()
            equation = equation.rstrip()
            if equation.endswith("$$"):
                equation = equation[:-2]
            equation_no += 1
            add_equation(doc, equation, styles, max(chapter, 1), equation_no)
            idx += 1
            continue

        if "|" in line and idx + 1 < len(lines) and TABLE_SEPARATOR_RE.match(lines[idx + 1]):
            rows = [parse_table_row(line)]
            idx += 2
            while idx < len(lines) and "|" in lines[idx] and lines[idx].strip():
                rows.append(parse_table_row(lines[idx]))
                idx += 1
            add_markdown_table(doc, rows, styles)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style=styles["body"])
            add_text_with_citations(p, bullet.group(1))
            p.paragraph_format.first_line_indent = None
            set_paragraph_numbering(p, numbering["bullet"])
            idx += 1
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style=styles["body"])
            add_text_with_citations(p, numbered.group(1))
            p.paragraph_format.first_line_indent = None
            set_paragraph_numbering(p, numbering["decimal"])
            idx += 1
            continue

        paragraph_lines = [line]
        idx += 1
        while idx < len(lines):
            candidate = lines[idx].strip()
            if not candidate:
                break
            if re.match(r"^(#{1,4})\s+", candidate) or FIGURE_RE.match(candidate) or candidate.startswith("$$"):
                break
            if candidate == "---PAGEBREAK---" or re.match(r"^[-*]\s+", candidate) or re.match(r"^\d+[.)]\s+", candidate):
                break
            if "|" in candidate and idx + 1 < len(lines) and TABLE_SEPARATOR_RE.match(lines[idx + 1]):
                break
            paragraph_lines.append(candidate)
            idx += 1
        p = doc.add_paragraph(style=styles["body"])
        add_text_with_citations(p, " ".join(paragraph_lines))


def build_draft(request: dict[str, Any], output: Path) -> None:
    metadata = request["metadata"]
    requirements = request.get("requirements") or {}
    content = request["content"]
    doc_type = request["document_type"]
    markdown = get_markdown(content)
    markdown, ordered_citations = citation_order(markdown, request.get("citations") or [])

    doc = Document()
    styles = configure_styles(doc)
    numbering = add_numbering(doc)
    for section_index, section in enumerate(doc.sections):
        configure_page(section, first_section=section_index == 0)
    add_cover(doc, metadata, doc_type)

    if doc_type == "hust_thesis":
        front = doc.add_section(WD_SECTION_START.NEW_PAGE)
        configure_page(front)
        set_page_numbering(front, 1, "upperRoman")
        configure_header_footer(front, "华中科技大学硕士学位论文")
        doc.add_paragraph("摘要", style=styles["unnumbered"])
        doc.add_paragraph(str(content["abstract_cn"]), style=styles["body"])
        add_keywords(doc, "关键词：", content["keywords_cn"], styles["body"])
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        doc.add_paragraph("ABSTRACT", style=styles["unnumbered"])
        doc.add_paragraph(str(content["abstract_en"]), style=styles["body"])
        add_keywords(doc, "Keywords: ", content["keywords_en"], styles["body"])
        if requirements.get("include_toc", True):
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            add_toc(doc, styles, markdown)
        body_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        configure_page(body_section)
        set_page_numbering(body_section, 1, None)
        body_section.header.is_linked_to_previous = True
        body_section.footer.is_linked_to_previous = True
    else:
        header = str(metadata.get("course") or "课程报告")
        if requirements.get("include_toc", True):
            front = doc.add_section(WD_SECTION_START.NEW_PAGE)
            configure_page(front)
            set_page_numbering(front, 1, "upperRoman")
            configure_header_footer(front, header)
            add_toc(doc, styles, markdown)
            body_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        else:
            body_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
        configure_page(body_section)
        set_page_numbering(body_section, 1, None)
        if requirements.get("include_toc", True):
            body_section.header.is_linked_to_previous = True
            body_section.footer.is_linked_to_previous = True
        else:
            configure_header_footer(body_section, header)

    render_markdown(doc, markdown, styles, numbering, request.get("assets") or [])
    add_references(doc, ordered_citations, styles)
    set_update_fields(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    apply_authoritative_template_parts(output)


def style_level(name: str, text: str) -> int | None:
    lowered = name.lower().replace(" ", "")
    for level in range(1, 5):
        if f"heading{level}" in lowered or f"标题{level}" in lowered or f"标题{['一','二','三','四'][level-1]}" in lowered:
            return level
    match = re.match(r"^(\d+(?:\.\d+){0,3})\s+", text)
    if match:
        return match.group(1).count(".") + 1
    return None


def normalize_runs(paragraph: Any, east_asia: str, latin: str, size: float) -> None:
    for run in paragraph.runs:
        bold, italic = run.bold, run.italic
        run_font(run, east_asia, latin, size)
        run.bold, run.italic = bold, italic
        rpr = run._element.get_or_add_rPr()
        for tag in ("w:color", "w:highlight", "w:shd", "w:u", "w:spacing"):
            remove_children(rpr, tag)


def clear_direct_typography(paragraph: Any) -> None:
    """Let the retained template style supply font family and size.

    Bold, italic, underline, superscript, field markup, and other semantic run
    properties are intentionally preserved.
    """
    for run in paragraph.runs:
        rpr = run._element.find(qn("w:rPr"))
        if rpr is None:
            continue
        for tag in ("w:rFonts", "w:sz", "w:szCs"):
            remove_children(rpr, tag)


def build_format(request: dict[str, Any], output: Path) -> None:
    source = Path(request["content"]["source_docx"])
    doc = Document(source)
    styles = configure_styles(doc)
    numbering = add_numbering(doc)
    requirements = request.get("requirements") or {}
    for section_index, section in enumerate(doc.sections):
        configure_page(section, first_section=section_index == 0)
        if requirements.get("replace_headers"):
            configure_header_footer(section, str(request["metadata"].get("course") or "课程报告"))

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        current = paragraph.style.name if paragraph.style is not None else ""
        level = style_level(current, text)
        if level:
            paragraph.style = doc.styles[f"Heading {level}"]
            remove_children(paragraph._p.get_or_add_pPr(), "w:numPr")
            clear_direct_typography(paragraph)
        elif current == styles["unnumbered"]:
            paragraph.style = doc.styles[styles["unnumbered"]]
            clear_direct_typography(paragraph)
        elif re.match(r"^(图|表)\s*\d", text) or current == styles["caption"]:
            paragraph.style = doc.styles[styles["caption"]]
            clear_direct_typography(paragraph)
        elif current in {"公式", "Equation"}:
            paragraph.style = doc.styles[styles["equation"]]
            clear_direct_typography(paragraph)
        elif text:
            paragraph.style = doc.styles[styles["body"]]
            clear_direct_typography(paragraph)

    for table in doc.tables:
        first_column = [row.cells[0].text.strip() for row in table.rows if row.cells]
        layout_labels = {"课程", "姓名", "学号", "教师", "日期", "学位申请人", "学科专业", "指导教师"}
        if first_column and sum(label in layout_labels for label in first_column) >= 2:
            continue
        table.autofit = False
        set_table_borders(table)
        cols = max((len(row.cells) for row in table.rows), default=1)
        width = 8730 // cols
        for ridx, row in enumerate(table.rows):
            for cell in row.cells:
                set_cell_width(cell, width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if ridx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    normalize_runs(paragraph, "宋体", "Times New Roman", 10.5)

    set_update_fields(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    apply_authoritative_template_parts(output)


def build(request_path: Path, output: Path) -> None:
    request = load_json(request_path)
    validate_request(request)
    if output.suffix.lower() != ".docx":
        raise BuildError("Output must end in .docx.")
    if output.exists() and request.get("content", {}).get("source_docx"):
        if output.resolve() == Path(request["content"]["source_docx"]).resolve():
            raise BuildError("Output must not overwrite the source DOCX.")
    if request["mode"] == "draft":
        build_draft(request, output)
    else:
        build_format(request, output)


def parse_docx(path: Path) -> tuple[zipfile.ZipFile, etree._Element, etree._Element]:
    if not path.is_file():
        raise BuildError(f"File not found: {path}")
    try:
        zf = zipfile.ZipFile(path)
        document = etree.fromstring(zf.read("word/document.xml"))
        styles = etree.fromstring(zf.read("word/styles.xml"))
    except (zipfile.BadZipFile, KeyError, etree.XMLSyntaxError) as exc:
        raise BuildError(f"Invalid DOCX package: {exc}") from exc
    return zf, document, styles


def style_map(styles: etree._Element) -> dict[str, str]:
    result: dict[str, str] = {}
    for style in styles.xpath("//w:style", namespaces=NS):
        sid = style.get(qn("w:styleId"), "")
        names = style.xpath("./w:name/@w:val", namespaces=NS)
        result[sid] = names[0] if names else sid
    return result


def find_style(styles: etree._Element, human_name: str) -> etree._Element | None:
    target = human_name.casefold()
    for style in styles.xpath("//w:style", namespaces=NS):
        names = style.xpath("./w:name/@w:val", namespaces=NS)
        if names and str(names[0]).casefold() == target:
            return style
    return None


def style_value(style: etree._Element, path: str, attr: str = "val") -> str | None:
    values = style.xpath(f"{path}/@w:{attr}", namespaces=NS)
    return str(values[0]) if values else None


def analyze_docx(path: Path) -> dict[str, Any]:
    zf, document, styles = parse_docx(path)
    names = set(zf.namelist())
    style_names = style_map(styles)
    used = Counter()
    for pstyle in document.xpath("//w:pPr/w:pStyle/@w:val", namespaces=NS):
        used[str(pstyle)] += 1
    fonts = Counter()
    for root in (styles, document):
        for rfonts in root.xpath("//w:rFonts", namespaces=NS):
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = rfonts.get(qn(f"w:{attr}"))
                if value:
                    fonts[value] += 1
    fields = [str(value).strip() for value in document.xpath("//w:instrText/text()", namespaces=NS) if str(value).strip()]
    sections = []
    for sect in document.xpath("//w:sectPr", namespaces=NS):
        pg = sect.find("w:pgSz", NS)
        mar = sect.find("w:pgMar", NS)
        sections.append(
            {
                "page": {etree.QName(k).localname: v for k, v in (pg.attrib.items() if pg is not None else [])},
                "margins": {etree.QName(k).localname: v for k, v in (mar.attrib.items() if mar is not None else [])},
            }
        )
    last = sections[-1] if sections else {"page": {}, "margins": {}}
    expected = {
        "page": {"w": "11906", "h": "16838"},
        "margins": {"top": "2552", "right": "1588", "bottom": "1588", "left": "1588"},
    }
    conflicts = []
    for key, expected_value in expected["page"].items():
        actual = last["page"].get(key)
        if actual and actual != expected_value:
            conflicts.append({"category": "page_geometry", "property": key, "external": actual, "academic_word": expected_value})
    for key, expected_value in expected["margins"].items():
        actual = last["margins"].get(key)
        if actual and actual != expected_value:
            conflicts.append({"category": "page_geometry", "property": f"margin_{key}", "external": actual, "academic_word": expected_value})
    nonportable = sorted(font for font in fonts if font not in ALLOWED_FONTS and not font.lower().startswith("minor"))
    if nonportable:
        conflicts.append({"category": "fonts", "external": nonportable, "academic_word": sorted(ALLOWED_FONTS)})
    return {
        "path": str(path.resolve()),
        "paragraphs": len(document.xpath("//w:p", namespaces=NS)),
        "tables": len(document.xpath("//w:tbl", namespaces=NS)),
        "sections": sections,
        "styles_defined": len(style_names),
        "styles_used": [{"id": sid, "name": style_names.get(sid, "undefined"), "count": count} for sid, count in used.most_common()],
        "fonts": dict(fonts.most_common()),
        "field_types": dict(Counter(field_type(value) for value in fields)),
        "has_addin_fields": any(field_type(value) == "ADDIN" for value in fields),
        "has_ole_embeddings": any(name.startswith("word/embeddings/") for name in names),
        "has_comments": "word/comments.xml" in names,
        "conflicts_with_fallback": conflicts,
    }


def field_type(instruction: str) -> str:
    stripped = instruction.strip()
    return re.split(r"[ \\]", stripped, maxsplit=1)[0].upper() if stripped else ""


def package_text(path: Path) -> str:
    zf, document, _ = parse_docx(path)
    del zf
    return "\n".join(text for text in document.xpath("//w:t/text()", namespaces=NS) if text)


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value)


def validate_docx(path: Path, request_path: Path | None = None, source: Path | None = None) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        zf, document, styles = parse_docx(path)
    except BuildError as exc:
        return {"ok": False, "errors": [str(exc)], "warnings": [], "metrics": {}}
    names = set(zf.namelist())
    bad_member = zf.testzip()
    if bad_member:
        errors.append(f"Corrupt ZIP member: {bad_member}")
    for name in names:
        if name.endswith(".xml") or name.endswith(".rels"):
            try:
                etree.fromstring(zf.read(name))
            except etree.XMLSyntaxError as exc:
                errors.append(f"Malformed XML in {name}: {exc}")

    visible_text = "\n".join(document.xpath("//w:t/text()", namespaces=NS))
    placeholder_hits = sorted(set(PLACEHOLDER_RE.findall(visible_text)))
    if placeholder_hits:
        errors.append(f"Unresolved placeholders: {placeholder_hits[:10]}")

    instructions = [str(value).strip() for value in document.xpath("//w:instrText/text()", namespaces=NS) if str(value).strip()]
    types = Counter(field_type(value) for value in instructions)
    unsupported_fields = sorted(ft for ft in types if ft and ft not in ALLOWED_FIELDS)
    if unsupported_fields:
        errors.append(f"Unsupported/plugin fields present: {unsupported_fields}")
    if any("NOTEEXPRESS" in value.upper() or "ENDNOTE" in value.upper() for value in instructions):
        errors.append("Reference-manager ADDIN fields remain in the document.")

    if any(name.startswith("word/embeddings/") for name in names):
        errors.append("Embedded OLE/Visio objects remain in word/embeddings.")
    if document.xpath("//w:object", namespaces=NS):
        errors.append("Legacy Word object elements remain in document.xml.")
    if "word/comments.xml" in names:
        errors.append("Word comments remain in the final document.")
    revisions = len(document.xpath("//w:ins | //w:del | //w:moveFrom | //w:moveTo", namespaces=NS))
    if revisions:
        errors.append(f"Tracked revisions remain: {revisions}")

    defined = set(style_map(styles))
    story_roots = [document]
    for story_name in sorted(names):
        if re.fullmatch(r"word/(?:header|footer)\d+\.xml", story_name) or story_name in {
            "word/footnotes.xml", "word/endnotes.xml"
        }:
            story_roots.append(etree.fromstring(zf.read(story_name)))
    referenced: set[str] = set()
    for story in story_roots:
        referenced.update(story.xpath("//w:pStyle/@w:val | //w:rStyle/@w:val | //w:tblStyle/@w:val", namespaces=NS))
    missing_styles = sorted(style for style in referenced if style not in defined)
    if missing_styles:
        errors.append(f"Referenced styles are undefined: {missing_styles}")

    numbering_ids: set[str] = set()
    if "word/numbering.xml" in names:
        numbering = etree.fromstring(zf.read("word/numbering.xml"))
        numbering_ids = set(numbering.xpath("//w:num/@w:numId", namespaces=NS))
    used_num_ids = set(document.xpath("//w:numPr/w:numId/@w:val", namespaces=NS))
    orphan_num = sorted(num for num in used_num_ids if num != "0" and num not in numbering_ids)
    if orphan_num:
        errors.append(f"Orphan numbering references: {orphan_num}")

    # Fidelity gate: these are not reconstructed equivalents. They must be the
    # retained DOTX parts byte-for-byte, and Word's competing style-effects
    # part must be absent.
    if "word/stylesWithEffects.xml" in names:
        errors.append("stylesWithEffects.xml remains and can override the authoritative DOTX styles in Microsoft Word.")
    if AUTHORITATIVE_DOTX.is_file():
        with zipfile.ZipFile(AUTHORITATIVE_DOTX, "r") as template:
            for part in (
                "word/styles.xml",
                "word/numbering.xml",
                "word/theme/theme1.xml",
                "word/fontTable.xml",
                "word/webSettings.xml",
            ):
                if part not in names or zf.read(part) != template.read(part):
                    errors.append(f"Authoritative template part differs from retained DOTX: {part}")
            if "word/settings.xml" not in names:
                errors.append("word/settings.xml is missing.")
            else:
                actual_settings = etree.fromstring(zf.read("word/settings.xml"))
                expected_settings = etree.fromstring(template.read("word/settings.xml"))
                for root in (actual_settings, expected_settings):
                    for node in root.xpath("./w:updateFields", namespaces=NS):
                        root.remove(node)
                if etree.tostring(actual_settings, method="c14n") != etree.tostring(expected_settings, method="c14n"):
                    errors.append("Document settings differ from retained DOTX beyond the required updateFields flag.")
    else:
        errors.append(f"Authoritative DOTX is missing: {AUTHORITATIVE_DOTX}")

    style_names = style_map(styles)
    expected_heading_ids = {"heading 1": "10", "heading 2": "2", "heading 3": "3", "heading 4": "4"}
    for paragraph in document.xpath("//w:p", namespaces=NS):
        pstyle = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
        if not pstyle:
            continue
        sid = str(pstyle[0])
        human_name = style_names.get(sid, sid).casefold()
        if human_name in expected_heading_ids:
            if sid != expected_heading_ids[human_name]:
                errors.append(f"{human_name} paragraph uses styleId {sid}; expected DOTX styleId {expected_heading_ids[human_name]}.")
            if paragraph.xpath("./w:pPr/w:numPr", namespaces=NS):
                errors.append(f"{human_name} paragraph contains direct numPr; the DOTX uses style-level numId 23 only.")

    unnumbered_style = find_style(styles, "无编号标题")
    unnumbered_num_id = style_value(unnumbered_style, "./w:pPr/w:numPr/w:numId") if unnumbered_style is not None else None
    if unnumbered_num_id != "0":
        errors.append(f"Unnumbered heading must use the DOTX numId 0; found {unnumbered_num_id!r}.")

    fonts = set()
    used_style_nodes = []
    critical_style_names = {
        "Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4",
        "宋体小四", "无编号标题", "题注", "图片居中", "公式", "参考文献",
        "TOC 1", "TOC 2", "TOC 3",
    }
    for style in styles.xpath("//w:style", namespaces=NS):
        sid = style.get(qn("w:styleId"), "")
        names_for_style = style.xpath("./w:name/@w:val", namespaces=NS)
        human_name = str(names_for_style[0]) if names_for_style else sid
        if sid in referenced or human_name in critical_style_names:
            used_style_nodes.append(style)
    font_roots = [*story_roots, *used_style_nodes]
    for root in font_roots:
        for rfonts in root.xpath(".//w:rFonts", namespaces=NS):
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = rfonts.get(qn(f"w:{attr}"))
                if value:
                    fonts.add(value)
    nonportable = sorted(font for font in fonts if font not in ALLOWED_FONTS)
    if nonportable:
        errors.append(f"Non-portable fonts present: {nonportable}")

    prior_level = 0
    for paragraph in document.xpath("//w:p", namespaces=NS):
        pstyle = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
        if not pstyle:
            continue
        name = style_names.get(str(pstyle[0]), str(pstyle[0]))
        match = re.search(r"heading\s*([1-4])", name, re.IGNORECASE)
        if match:
            level = int(match.group(1))
            if prior_level and level > prior_level + 1:
                errors.append(f"Heading hierarchy skips from level {prior_level} to {level}.")
                break
            prior_level = level

    for paragraph in document.xpath("//w:p[w:pPr/w:pStyle[@w:val='aff1']]", namespaces=NS):
        for run in paragraph.xpath("./w:r", namespaces=NS):
            run_text = "".join(run.xpath(".//w:t/text()", namespaces=NS))
            if VISIBLE_CITATION_RE.search(run_text):
                continue
            if run.xpath("./w:rPr/w:rFonts | ./w:rPr/w:sz | ./w:rPr/w:szCs", namespaces=NS):
                errors.append(f"Body run contains direct font/size overrides instead of retained 宋体小四 style: {run_text[:40]}")
                break

    sections = document.xpath("//w:sectPr", namespaces=NS)
    expected_margins = {
        "top": "2552", "right": "1588", "bottom": "1588", "left": "1588",
        "header": "851", "footer": "964", "gutter": "0",
    }
    for index, section in enumerate(sections):
        pg_sz = section.find("w:pgSz", NS)
        pg_mar = section.find("w:pgMar", NS)
        cols = section.find("w:cols", NS)
        grid = section.find("w:docGrid", NS)
        if pg_sz is None or pg_sz.get(qn("w:w")) != "11906" or pg_sz.get(qn("w:h")) != "16838":
            errors.append(f"Section {index + 1} page size differs from the retained DOTX A4 geometry.")
        expected_code = "9" if index == 0 else None
        if pg_sz is not None and pg_sz.get(qn("w:code")) != expected_code:
            errors.append(f"Section {index + 1} page-size code differs from retained DOTX: expected {expected_code!r}.")
        actual_margins = {} if pg_mar is None else {
            key: pg_mar.get(qn(f"w:{key}")) for key in expected_margins
        }
        if actual_margins != expected_margins:
            errors.append(f"Section {index + 1} margins differ from retained DOTX: {actual_margins}.")
        expected_col_space = "425" if index == 0 else "720"
        if cols is None or cols.get(qn("w:space")) != expected_col_space:
            errors.append(f"Section {index + 1} column spacing must be {expected_col_space} DXA.")
        if grid is None or grid.get(qn("w:type")) != "linesAndChars" or grid.get(qn("w:linePitch")) != "312":
            errors.append(f"Section {index + 1} document grid must be linesAndChars with 312-DXA line pitch.")

    # Header/footer stories must use the retained styles rather than visually
    # similar direct formatting.
    for story_name in sorted(names):
        if not re.fullmatch(r"word/header\d+\.xml", story_name):
            continue
        story = etree.fromstring(zf.read(story_name))
        if not "".join(story.xpath("//w:t/text()", namespaces=NS)).strip():
            continue
        top_paragraphs = story.xpath("./w:p", namespaces=NS)
        if not top_paragraphs:
            errors.append(f"Non-empty {story_name} has no top-level paragraph.")
            continue
        paragraph = top_paragraphs[0]
        if paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS) != ["ad"]:
            errors.append(f"{story_name} must use retained header styleId ad.")
        if paragraph.xpath("./w:pPr/w:pBdr | ./w:pPr/w:jc", namespaces=NS):
            errors.append(f"{story_name} contains direct header border/alignment overrides.")
    for story_name in sorted(names):
        if not re.fullmatch(r"word/footer\d+\.xml", story_name):
            continue
        story = etree.fromstring(zf.read(story_name))
        if not story.xpath("//w:instrText[contains(normalize-space(.), 'PAGE')]", namespaces=NS):
            continue
        top_paragraphs = story.xpath("./w:p", namespaces=NS)
        paragraph = top_paragraphs[0] if top_paragraphs else None
        if paragraph is None or paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS) != ["af1"]:
            errors.append(f"{story_name} must use retained footer styleId af1.")
            continue
        frame = paragraph.find("w:pPr/w:framePr", NS)
        expected_frame = {"wrap": "around", "vAnchor": "text", "hAnchor": "margin", "xAlign": "center", "y": "1"}
        actual_frame = {} if frame is None else {key: frame.get(qn(f"w:{key}")) for key in expected_frame}
        if actual_frame != expected_frame:
            errors.append(f"{story_name} page number frame differs from retained DOTX: {actual_frame}.")
        for run in paragraph.xpath("./w:r", namespaces=NS):
            if run.xpath(".//w:fldChar | .//w:instrText | .//w:t", namespaces=NS):
                styles_for_run = run.xpath("./w:rPr/w:rStyle/@w:val", namespaces=NS)
                if styles_for_run != ["af2"]:
                    errors.append(f"{story_name} PAGE field run does not use retained page-number styleId af2.")
                    break

    # Check internal relationship targets for the main document.
    rel_name = "word/_rels/document.xml.rels"
    if rel_name in names:
        rels = etree.fromstring(zf.read(rel_name))
        for rel in rels.xpath("/*[local-name()='Relationships']/*[local-name()='Relationship']"):
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            resolved = posixpath.normpath(posixpath.join("word", target))
            if resolved not in names:
                errors.append(f"Broken relationship target: {target}")

        # The retained settings reference separator IDs in both note stories.
        # Word treats a missing companion part/relationship/override as package
        # corruption even when there are no user-created footnotes or endnotes.
        settings_for_notes = etree.fromstring(zf.read("word/settings.xml")) if "word/settings.xml" in names else None
        content_types_for_notes = etree.fromstring(zf.read("[Content_Types].xml"))
        for note_name, setting_tag, rel_suffix, root_tag in (
            ("footnotes", "footnotePr", "/footnotes", "footnotes"),
            ("endnotes", "endnotePr", "/endnotes", "endnotes"),
        ):
            if settings_for_notes is None or not settings_for_notes.xpath(f"./w:{setting_tag}", namespaces=NS):
                continue
            part_name = f"word/{note_name}.xml"
            if part_name not in names:
                errors.append(f"{part_name} is required by settings.xml and is missing (Word repair trigger).")
                continue
            if not any(str(rel.get("Type", "")).endswith(rel_suffix) and rel.get("Target") == f"{note_name}.xml" for rel in rels):
                errors.append(f"Document relationship for {part_name} is missing (Word repair trigger).")
            override_name = f"/{part_name}"
            if not any(node.get("PartName") == override_name for node in content_types_for_notes):
                errors.append(f"Content-type override for {part_name} is missing (Word repair trigger).")
            note_root = etree.fromstring(zf.read(part_name))
            if etree.QName(note_root).localname != root_tag:
                errors.append(f"{part_name} has the wrong root element.")
            ids = set(note_root.xpath("./w:*[self::w:footnote or self::w:endnote]/@w:id", namespaces=NS))
            if not {"-1", "0"}.issubset(ids):
                errors.append(f"{part_name} lacks required separator IDs -1 and 0.")

    # TOC cached results must preserve the template's mixed-script sizing.
    # The template fixes only w:sz=21 on cached entries; w:szCs is inherited
    # as 24 (12 pt), so numeric prefixes are not undersized.
    if document.xpath("//w:p[w:pPr/w:pStyle[@w:val='TOC3']]", namespaces=NS):
        errors.append("TOC contains a level-3 entry; only Heading 1 and Heading 2 may appear in the table of contents.")
    toc_instructions = [str(value) for value in document.xpath("//w:instrText[starts-with(normalize-space(.), 'TOC ')]/text()", namespaces=NS)]
    if toc_instructions and not all('\\o "1-2"' in value and '\\o "1-3"' not in value for value in toc_instructions):
        errors.append('TOC field must use the outline range \\o "1-2".')
    toc_paragraphs = document.xpath(
        "//w:p[w:pPr/w:pStyle[@w:val='TOC1' or @w:val='TOC2']]",
        namespaces=NS,
    )
    for paragraph in toc_paragraphs:
        if paragraph.xpath(".//w:rPr/w:szCs", namespaces=NS):
            errors.append("TOC cached result contains a direct szCs override; numeric/Latin text must inherit 12 pt from the retained template style.")
            break

    # Every content table must have a real, repeatable header row.  Cover
    # metadata tables are layout furniture and are intentionally excluded.
    layout_labels = {"课程", "姓名", "学号", "教师", "日期", "学位申请人", "学科专业", "指导教师"}
    content_table_count = 0
    repeated_header_count = 0
    body_nodes = document.xpath("//w:body", namespaces=NS)
    body_node = body_nodes[0] if body_nodes else None
    first_section_break_index = None
    if body_node is not None:
        for child_index, child in enumerate(body_node):
            if etree.QName(child).localname == "p" and child.xpath("./w:pPr/w:sectPr", namespaces=NS):
                first_section_break_index = child_index
                break
    for table in document.xpath("//w:tbl", namespaces=NS):
        if (
            body_node is not None
            and table.getparent() is body_node
            and first_section_break_index is not None
            and body_node.index(table) < first_section_break_index
        ):
            # Cover-page metadata tables are layout furniture, never data tables.
            continue
        first_column = []
        for row in table.xpath("./w:tr", namespaces=NS):
            cells = row.xpath("./w:tc", namespaces=NS)
            first_column.append("".join(cells[0].xpath(".//w:t/text()", namespaces=NS)).strip() if cells else "")
        if first_column and sum(label in layout_labels for label in first_column) >= 2:
            continue
        content_table_count += 1
        first_rows = table.xpath("./w:tr[1]", namespaces=NS)
        if first_rows and first_rows[0].xpath("./w:trPr/w:tblHeader[@w:val='true' or @w:val='1']", namespaces=NS):
            repeated_header_count += 1
        else:
            errors.append("Content table is missing a real repeating header row (w:tblHeader).")

    # Figures must be inline drawings followed by a genuine caption paragraph:
    # retained caption style + visible STYLEREF/SEQ fields, never a text box or
    # typed chapter/sequence number.
    if document.xpath("//w:txbxContent", namespaces=NS):
        errors.append("Text-box content is present; figures and captions must use inline paragraphs.")
    figure_count = 0
    caption_field_count = 0
    for paragraph in document.xpath("//w:body//w:p[.//w:drawing or .//w:pict]", namespaces=NS):
        figure_count += 1
        sibling = paragraph.getnext()
        while sibling is not None and etree.QName(sibling).localname not in {"p", "tbl"}:
            sibling = sibling.getnext()
        if sibling is None or etree.QName(sibling).localname != "p" or sibling.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS) != ["aff3"]:
            errors.append("Inline figure is not immediately followed by retained caption styleId aff3.")
            continue
        instructions = [str(value).strip() for value in sibling.xpath(".//w:instrText/text()", namespaces=NS)]
        if not any(value.startswith("STYLEREF 1") for value in instructions) or not any(value.startswith("SEQ 图") for value in instructions):
            errors.append("Figure caption must contain real STYLEREF and SEQ 图 fields.")
            continue
        field_depth = 0
        outside_text = []
        for run in sibling.xpath("./w:r", namespaces=NS):
            for child in run:
                if child.tag == qn("w:fldChar") and child.get(qn("w:fldCharType")) == "begin":
                    field_depth += 1
                elif child.tag == qn("w:fldChar") and child.get(qn("w:fldCharType")) == "end":
                    field_depth = max(0, field_depth - 1)
                elif child.tag == qn("w:t") and field_depth == 0:
                    outside_text.append(child.text or "")
        visible_outside = "".join(outside_text)
        prefix = visible_outside.split("-", 1)[0]
        if re.search(r"\d", prefix):
            errors.append("Figure caption contains a typed chapter/sequence number outside Word fields.")
            continue
        caption_field_count += 1

    if source:
        if not source.is_file():
            errors.append(f"Source for text preservation does not exist: {source}")
        else:
            source_text = normalize_text(package_text(source))
            output_text = normalize_text(visible_text)
            if source_text != output_text:
                errors.append("Format-mode body text differs from the source document.")

    if request_path:
        request = load_json(request_path)
        if request.get("mode") == "draft":
            for marker in CITATION_RE.findall(visible_text):
                errors.append(f"Unresolved citation marker remains: {marker}")
            for citation in request.get("citations") or []:
                if citation.get("id") and citation.get("title") and f"[@{citation['id']}]" in get_markdown(request.get("content") or {}):
                    if str(citation["title"]) not in visible_text:
                        errors.append(f"Cited reference missing from bibliography: {citation['id']}")

            expected_citation_numbers = set(range(1, len(citation_order(get_markdown(request.get("content") or {}), request.get("citations") or [])[1]) + 1))
            superscript_numbers: set[int] = set()
            in_bibliography = False
            for paragraph in document.xpath("//w:p", namespaces=NS):
                paragraph_text = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)).strip()
                if paragraph_text == "参考文献":
                    in_bibliography = True
                    continue
                pstyle = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
                style_name = style_names.get(str(pstyle[0]), str(pstyle[0])) if pstyle else ""
                if in_bibliography or style_name == "参考文献":
                    continue
                for run in paragraph.xpath(".//w:r", namespaces=NS):
                    run_text = "".join(run.xpath(".//w:t/text()", namespaces=NS))
                    numbers = {int(value) for value in VISIBLE_CITATION_RE.findall(run_text)}
                    if not numbers:
                        continue
                    vert = run.xpath("./w:rPr/w:vertAlign/@w:val", namespaces=NS)
                    if not vert or str(vert[0]) != "superscript":
                        errors.append(f"In-text citation is not superscript: {run_text}")
                    else:
                        superscript_numbers.update(numbers)
                    rpr = run.find("w:rPr", NS)
                    color = None if rpr is None or rpr.find("w:color", NS) is None else rpr.find("w:color", NS).get(qn("w:val"))
                    kern = None if rpr is None or rpr.find("w:kern", NS) is None else rpr.find("w:kern", NS).get(qn("w:val"))
                    size = None if rpr is None or rpr.find("w:sz", NS) is None else rpr.find("w:sz", NS).get(qn("w:val"))
                    if (color, kern, size) != ("080000", "0", "24"):
                        errors.append(f"In-text citation formatting differs from retained template: {run_text}")
            if not expected_citation_numbers.issubset(superscript_numbers):
                missing = sorted(expected_citation_numbers - superscript_numbers)
                errors.append(f"Citations missing superscript in body text: {missing}")

            page_zones = []
            for section in sections:
                pg_type = section.find("w:pgNumType", NS)
                page_zones.append({
                    "fmt": pg_type.get(qn("w:fmt")) if pg_type is not None else None,
                    "start": pg_type.get(qn("w:start")) if pg_type is not None else None,
                })
            include_toc = bool((request.get("requirements") or {}).get("include_toc", True))
            if not page_zones or page_zones[0] != {"fmt": "upperRoman", "start": "1"}:
                errors.append(f"Cover section must retain hidden Roman numbering from I; page zones are {page_zones}.")
            if request.get("document_type") == "hust_thesis" or include_toc:
                if len(page_zones) < 3 or page_zones[-2] != {"fmt": "upperRoman", "start": "1"}:
                    errors.append(f"Front matter must use uppercase Roman page numbers from I; page zones are {page_zones}.")
            if len(page_zones) < 2 or page_zones[-1] != {"fmt": None, "start": "1"}:
                errors.append(f"Body must restart Arabic page numbering at 1; page zones are {page_zones}.")
            if len(sections) >= 3:
                third = sections[-1]
                if third.xpath("./w:headerReference | ./w:footerReference", namespaces=NS):
                    errors.append("Body section must link to the previous header/footer exactly like the retained template.")

    update_fields = False
    if "word/settings.xml" in names:
        settings = etree.fromstring(zf.read("word/settings.xml"))
        update_fields = bool(settings.xpath("//w:updateFields[@w:val='true' or @w:val='1']", namespaces=NS))
    if not update_fields:
        errors.append("w:updateFields is not enabled.")

    metrics = {
        "paragraphs": len(document.xpath("//w:p", namespaces=NS)),
        "tables": len(document.xpath("//w:tbl", namespaces=NS)),
        "sections": len(document.xpath("//w:sectPr", namespaces=NS)),
        "fields": dict(types),
        "fonts": sorted(fonts),
        "revisions": revisions,
        "placeholders": len(placeholder_hits),
        "content_tables": content_table_count,
        "repeating_table_headers": repeated_header_count,
        "figures": figure_count,
        "field_captions": caption_field_count,
    }
    return {"ok": not errors, "errors": errors, "warnings": warnings, "metrics": metrics}


def make_masters(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_request = {
        "mode": "draft",
        "document_type": "course_report",
        "metadata": {
            "title": "课程报告排版示例",
            "author": "Academic Word",
            "course": "课程报告",
            "date": "2026年",
        },
        "requirements": {"include_toc": True},
        "content": {"markdown": "# 正文结构\n本母版展示正文、标题、图表、公式和参考文献所使用的统一样式。\n## 二级标题\n正文采用宋体小四，英文与数字采用 Times New Roman。"},
        "citations": [],
        "assets": [],
    }
    thesis_request = {
        "mode": "draft",
        "document_type": "hust_thesis",
        "metadata": {
            "title": "学位论文排版示例",
            "title_en": "Thesis Formatting Example",
            "author": "Academic Word",
            "author_en": "Academic Word",
            "student_id": "00000000",
            "major": "排版规范",
            "major_en": "Document Formatting",
            "advisor": "Academic Word",
            "advisor_en": "Academic Word",
            "degree_en": "Master's Degree",
            "institution": "华中科技大学",
            "date": "2026年",
            "date_en": "2026",
            "classification": "示例",
            "secrecy": "公开",
        },
        "requirements": {"include_toc": True},
        "content": {
            "abstract_cn": "本母版仅用于保存学位论文结构和样式，不包含提交内容。",
            "keywords_cn": ["论文格式", "Word"],
            "abstract_en": "This master preserves the structure and styles used for thesis production.",
            "keywords_en": ["thesis format", "Word"],
            "markdown": "# 正文结构\n本母版展示学位论文正文所使用的标题、正文和分页样式。\n## 二级标题\n正式文档由完整资料生成。",
        },
        "citations": [],
        "assets": [],
    }
    temp_report = output_dir / ".clean-report-request.json"
    temp_thesis = output_dir / ".clean-thesis-request.json"
    save_json(temp_report, report_request)
    save_json(temp_thesis, thesis_request)
    build(temp_report, output_dir / "clean-report.docx")
    build(temp_thesis, output_dir / "clean-thesis.docx")
    temp_report.unlink()
    temp_thesis.unlink()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    p_build = sub.add_parser("build", help="Build a DOCX from a normalized request")
    p_build.add_argument("--request", type=Path, required=True)
    p_build.add_argument("--output", type=Path, required=True)

    p_analyze = sub.add_parser("analyze", help="Analyze a DOCX/DOTX and report style conflicts")
    p_analyze.add_argument("--input", type=Path, required=True)
    p_analyze.add_argument("--json-report", type=Path)

    p_validate = sub.add_parser("validate", help="Run the final structural gate")
    p_validate.add_argument("--input", type=Path, required=True)
    p_validate.add_argument("--request", type=Path)
    p_validate.add_argument("--source", type=Path)
    p_validate.add_argument("--json-report", type=Path)

    p_masters = sub.add_parser("make-masters", help="Generate the two clean bundled masters")
    p_masters.add_argument("--output-dir", type=Path, required=True)

    p_toc = sub.add_parser("cache-toc", help="Materialize verified page numbers in the cached TOC result")
    p_toc.add_argument("--input", type=Path, required=True)
    p_toc.add_argument("--page-map", type=Path, required=True)
    p_toc.add_argument("--output", type=Path, required=True)

    args = parser.parse_args()
    try:
        if args.command == "build":
            build(args.request, args.output)
            print(json.dumps({"ok": True, "output": str(args.output.resolve())}, ensure_ascii=False))
            return 0
        if args.command == "cache-toc":
            cache_toc_pages(args.input, args.page_map, args.output)
            print(json.dumps({"ok": True, "output": str(args.output.resolve())}, ensure_ascii=False))
            return 0
        if args.command == "analyze":
            result = analyze_docx(args.input)
        elif args.command == "validate":
            result = validate_docx(args.input, args.request, args.source)
        else:
            make_masters(args.output_dir)
            print(json.dumps({"ok": True, "output_dir": str(args.output_dir.resolve())}, ensure_ascii=False))
            return 0
        if getattr(args, "json_report", None):
            save_json(args.json_report, result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("ok", True) else 1
    except (BuildError, OSError, ValueError, KeyError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
